"""
generate_slides_v12.py
──────────────────────────────────────────────────────────────────────────
高效相参合成 · 研究背景与需求分析（单页 PPT）

布局（16:9  13.33" × 7.50"）：
  ┌──────────────────────────────────────────────────────────────────┐
  │  标题栏                                                            │
  ├──────────────────────────────────────────────────────────────────┤
  │  [研究目的]  全宽横幅 + 三项改进方向关键词徽章                       │
  ├─────────────────────────────┬────────────────────────────────────┤
  │  系统需求（左上卡片）          │                                    │
  │  · 有限通信带宽/计算资源       │  现有方法不足（右列全高，红/橙警示）  │
  │  · 实时相参合成               │  ① 计算复杂度高                    │
  │  · 工程部署灵活性             │  ② 同步精度要求苛刻                 │
  ├─────────────────────────────┤  ③ 非理想条件鲁棒性差               │
  │  性能需求（左下卡片）          │  ④ 闭环反馈—高动态难适应           │
  │  · 逼近理论最优增益           │  ⑤ 开环方法收敛受限                 │
  │  · 反隐身/弱目标/抗干扰       │                                    │
  │  · 高检测/跟踪精度            │                                    │
  ├─────────────────────────────┴────────────────────────────────────┤
  │  底部横条：三项改进技术方向                                          │
  ├──────────────────────────────────────────────────────────────────┤
  │  页脚                                                              │
  └──────────────────────────────────────────────────────────────────┘

风格：与 generate_slides_v10/v11.py 一致，布局比v11更丰富。

运行：
    python generate_slides_v12.py
输出：
    coherent_synthesis_slides_v12.pptx
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# ━━━━━━━━━━━━━━━━━━━━━━ 调色板（与 v10/v11 一致）━━━━━━━━━━━━━━━━━━━━━━
BG        = RGBColor(0xF4, 0xF7, 0xFC)
TITLE_BAR = RGBColor(0x1F, 0x38, 0x64)
TITLE_BAR2 = RGBColor(0x2E, 0x74, 0xB5)
DARK_TEXT  = RGBColor(0x1A, 0x1A, 0x3A)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
FOOTER_BG  = RGBColor(0xE8, 0xF0, 0xFB)
COL_NAVY   = RGBColor(0x0D, 0x2A, 0x6E)

# 研究目的横幅
OBJ_BAR  = RGBColor(0x1A, 0x3A, 0x6E)
OBJ_BAR2 = RGBColor(0x12, 0x2B, 0x55)

# 系统需求 — 蓝绿色系
SYS_HDR = RGBColor(0x0E, 0x55, 0x8C)
SYS_BG  = RGBColor(0xE5, 0xF0, 0xFB)
SYS_ACC = RGBColor(0x2E, 0x74, 0xB5)

# 性能需求 — 深绿色系
PER_HDR = RGBColor(0x15, 0x60, 0x3A)
PER_BG  = RGBColor(0xE5, 0xF5, 0xED)
PER_ACC = RGBColor(0x1A, 0x7A, 0x45)

# 现有方法不足 — 橙红警示色系
GAP_HDR = RGBColor(0x8B, 0x22, 0x00)
GAP_BG  = RGBColor(0xFD, 0xF0, 0xEB)
GAP_ACC = RGBColor(0xC0, 0x38, 0x00)

# 底部横条
BOT_COL = RGBColor(0x0D, 0x2A, 0x6E)

# 研究目的关键词徽章颜色
BADGE_COLS = [
    RGBColor(0x2E, 0x74, 0xB5),
    RGBColor(0x0E, 0x74, 0x64),
    RGBColor(0xB5, 0x55, 0x00),
]

# ━━━━━━━━━━━━━━━━━━━━━━ 尺寸常量 ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
SW      = Inches(13.33)
SH      = Inches(7.50)
MARGIN  = Inches(0.32)
TITLE_H = Inches(0.82)
CT      = TITLE_H + Inches(0.12)
FTR_H   = Inches(0.46)
FOOTER_Y = SH - FTR_H
CW      = SW - 2 * MARGIN


# ━━━━━━━━━━━━━━━━━━━━━━ 工具函数（与 v10/v11 一致）━━━━━━━━━━━━━━━━━━━

def set_bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def rect(slide, l, t, w, h, fill, lc=None, lw=Pt(0.75), rnd=False):
    s = slide.shapes.add_shape(5 if rnd else 1, l, t, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if lc:
        s.line.color.rgb = lc
        s.line.width = lw
    else:
        s.line.fill.background()
    return s


def hline(slide, x1, y, x2, color, width=Pt(0.6)):
    c = slide.shapes.add_connector(1, x1, y, x2, y)
    c.line.color.rgb = color
    c.line.width = width


def vline(slide, x, y1, y2, color, width=Pt(0.6)):
    c = slide.shapes.add_connector(1, x, y1, x, y2)
    c.line.color.rgb = color
    c.line.width = width


def tb(slide, text, l, t, w, h,
       sz=10, bold=False, color=DARK_TEXT,
       align=PP_ALIGN.LEFT, italic=False, wrap=True):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf  = box.text_frame
    tf.word_wrap = wrap
    p   = tf.paragraphs[0]
    p.alignment = align
    r   = p.add_run()
    r.text = text
    r.font.size   = Pt(sz)
    r.font.bold   = bold
    r.font.italic = italic
    r.font.color.rgb = color
    return box


def circle_badge(slide, cx, cy, r_inch, fill, text,
                 sz=10, bold=True, tc=WHITE):
    d = Inches(r_inch * 2)
    s = slide.shapes.add_shape(9,
                                cx - Inches(r_inch),
                                cy - Inches(r_inch), d, d)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    tf = s.text_frame
    tf.word_wrap = False
    p  = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r  = p.add_run()
    r.text = text
    r.font.size  = Pt(sz)
    r.font.bold  = bold
    r.font.color.rgb = tc


def title_bar(slide, title_main, title_sub, page, total=1):
    rect(slide, Inches(0), Inches(0), Inches(0.18), TITLE_H, COL_NAVY)
    rect(slide, Inches(0.18), Inches(0), SW - Inches(0.18), TITLE_H, TITLE_BAR)
    hline(slide, Inches(0), TITLE_H, SW, TITLE_BAR2, Pt(2.5))
    tb(slide, title_main,
       Inches(0.36), Inches(0.10), Inches(10.5), Inches(0.38),
       sz=20, bold=True, color=WHITE)
    tb(slide, title_sub,
       Inches(0.36), Inches(0.50), Inches(10.5), Inches(0.26),
       sz=10.5, color=RGBColor(0xB0, 0xCC, 0xF0))
    circle_badge(slide, SW - Inches(0.52), TITLE_H / 2, 0.28,
                 TITLE_BAR2, f"{page}/{total}", sz=11)


def footer(slide, text):
    rect(slide, Inches(0), FOOTER_Y, SW, FTR_H, FOOTER_BG)
    hline(slide, MARGIN, FOOTER_Y, SW - MARGIN, TITLE_BAR2, Pt(1.5))
    rect(slide, MARGIN, FOOTER_Y + Inches(0.10),
         Inches(0.05), Inches(0.28), TITLE_BAR2)
    tb(slide, "  " + text,
       MARGIN + Inches(0.10), FOOTER_Y + Inches(0.09),
       CW - Inches(0.14), Inches(0.34),
       sz=8.8, italic=True, color=COL_NAVY)


# ━━━━━━━━━━━━━━━━━━━━━━━━━━ 专用区块构建函数 ━━━━━━━━━━━━━━━━━━━━━━━━━

def left_card(slide, l, t, w, h,
              hdr_col, bg_col, acc_col,
              icon_char, title_text, subtitle_text,
              bullets, hdr_h=Inches(0.40),
              title_sz=11, body_sz=8.8):
    """
    左侧两个卡片（系统需求 / 性能需求）的通用构建函数。
    带彩色头带、左侧强调色竖条、彩色图标圆。
    bullets: [(text, is_accent), ...]
    """
    ACC_W = Inches(0.10)
    # 阴影
    rect(slide, l + Inches(0.022), t + Inches(0.022), w, h,
         RGBColor(0xC0, 0xC8, 0xD8))
    # 主背景
    rect(slide, l, t, w, h, bg_col, lc=hdr_col, lw=Pt(0.9))
    # 头带
    rect(slide, l, t, w, hdr_h, hdr_col)
    # 左侧竖条（正文区）
    rect(slide, l, t + hdr_h, ACC_W, h - hdr_h, hdr_col)
    # 图标圆
    circle_badge(slide, l + Inches(0.26), t + hdr_h / 2,
                 0.165, WHITE, icon_char, sz=13, tc=hdr_col, bold=True)
    # 标题
    bx = l + Inches(0.55)
    bw = w - bx + l - Inches(0.08)
    tb(slide, title_text, bx, t + Inches(0.05),
       bw, Inches(0.24), sz=title_sz, bold=True, color=WHITE)
    if subtitle_text:
        tb(slide, subtitle_text, bx, t + hdr_h - Inches(0.12),
           bw, Inches(0.14), sz=8, color=RGBColor(0xC0, 0xE8, 0xFF))
    # 正文 bullets
    pad_l = l + Inches(0.18)
    pad_w = w - Inches(0.24)
    yy = t + hdr_h + Inches(0.07)
    lh = Inches(0.195)
    gp = Inches(0.012)
    for (txt, accent) in bullets:
        clr = acc_col if accent else DARK_TEXT
        tb(slide, txt, pad_l, yy, pad_w, lh,
           sz=body_sz, color=clr, bold=accent)
        yy += lh + gp


def gap_card(slide, l, t, w, h, items, hdr_h=Inches(0.42)):
    """
    右侧"现有方法不足"卡片。
    items: [(num_str, title, detail), ...] 每条pain point
    """
    ACC_W = Inches(0.10)
    # 阴影
    rect(slide, l + Inches(0.022), t + Inches(0.022), w, h,
         RGBColor(0xC0, 0xC8, 0xD8))
    # 主背景
    rect(slide, l, t, w, h, GAP_BG, lc=GAP_HDR, lw=Pt(0.9))
    # 头带
    rect(slide, l, t, w, hdr_h, GAP_HDR)
    # 左侧竖条（正文区）
    rect(slide, l, t + hdr_h, ACC_W, h - hdr_h, GAP_HDR)
    # 头带图标圆
    circle_badge(slide, l + Inches(0.26), t + hdr_h / 2,
                 0.165, WHITE, "⚠", sz=12, tc=GAP_HDR, bold=True)
    # 头带标题
    bx = l + Inches(0.55)
    bw = w - bx + l - Inches(0.08)
    tb(slide, "现有方法不足", bx, t + Inches(0.06),
       bw, Inches(0.22), sz=11.5, bold=True, color=WHITE)
    tb(slide, "Limitations of Existing Coherent Synthesis Methods",
       bx, t + hdr_h - Inches(0.13),
       bw, Inches(0.14), sz=7.8,
       color=RGBColor(0xFF, 0xCC, 0xAA))

    # Pain-point strips
    pad_l   = l + Inches(0.16)
    pad_w   = w - Inches(0.22)
    item_h  = (h - hdr_h - Inches(0.10)) / len(items) - Inches(0.022)
    yy      = t + hdr_h + Inches(0.05)
    strip_gap = Inches(0.022)

    # alternating strip colors
    STRIP_A = RGBColor(0xFF, 0xE8, 0xDE)
    STRIP_B = RGBColor(0xFD, 0xF5, 0xF2)

    for i, (num_s, title, detail) in enumerate(items):
        sy = yy + i * (item_h + strip_gap)
        # strip bg
        sc = STRIP_A if i % 2 == 0 else STRIP_B
        rect(slide, pad_l, sy, pad_w, item_h, sc,
             lc=RGBColor(0xE8, 0xB0, 0x90), lw=Pt(0.5))
        # num badge
        circle_badge(slide, pad_l + Inches(0.175), sy + item_h / 2,
                     0.145, GAP_HDR, num_s, sz=8.5, tc=WHITE)
        # title
        tb(slide, title,
           pad_l + Inches(0.40), sy + Inches(0.03),
           pad_w - Inches(0.44), Inches(0.19),
           sz=9, bold=True, color=GAP_HDR)
        # detail
        tb(slide, detail,
           pad_l + Inches(0.40), sy + Inches(0.20),
           pad_w - Inches(0.44), item_h - Inches(0.22),
           sz=8, color=DARK_TEXT)


# ━━━━━━━━━━━━━━━━━━━━━━━━━━ MAIN ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━

prs = Presentation()
prs.slide_width  = SW
prs.slide_height = SH
blank = prs.slide_layouts[6]

s = prs.slides.add_slide(blank)
set_bg(s, BG)

# ── 1. 标题栏 ──────────────────────────────────────────────────────
title_bar(s,
          "高效相参合成  ·  研究背景与需求分析",
          "Research Background & Requirements: Efficient Coherent Synthesis "
          "under Resource Constraints and Real-Time Battlefield Demands",
          1)

# ── 2. 全局布局常量 ─────────────────────────────────────────────────
OBJ_H       = Inches(0.76)   # 研究目的横幅高度
BOT_STRIP_H = Inches(0.38)
BOT_Y       = FOOTER_Y - BOT_STRIP_H - Inches(0.04)

OBJ_Y    = CT
BODY_Y   = OBJ_Y + OBJ_H + Inches(0.10)
BODY_H   = BOT_Y - BODY_Y - Inches(0.04)

# 左列（系统需求 + 性能需求）、右列（现有方法不足）
LEFT_W   = Inches(6.20)
GAP_COL  = Inches(0.14)
RIGHT_W  = CW - LEFT_W - GAP_COL

LEFT_X   = MARGIN
RIGHT_X  = LEFT_X + LEFT_W + GAP_COL

# 左列两行
LEFT_GAP  = Inches(0.10)
SYS_H     = Inches(2.42)     # 系统需求卡片高度
PER_H     = BODY_H - SYS_H - LEFT_GAP  # 性能需求卡片高度（余量）

SYS_Y = BODY_Y
PER_Y = SYS_Y + SYS_H + LEFT_GAP

# ── 3. 研究目的 全宽横幅 ─────────────────────────────────────────────
rect(s, MARGIN + Inches(0.022), OBJ_Y + Inches(0.022),
     CW, OBJ_H, RGBColor(0xA0, 0xAA, 0xCC))
rect(s, MARGIN, OBJ_Y, CW, OBJ_H, OBJ_BAR2)
# 左侧"目的"标签
rect(s, MARGIN, OBJ_Y, Inches(0.76), OBJ_H,
     RGBColor(0x2E, 0x74, 0xB5))
tb(s, "研究\n目的",
   MARGIN + Inches(0.06), OBJ_Y + Inches(0.08),
   Inches(0.64), Inches(0.56),
   sz=10, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
# 目的正文
tb(s,
   "满足系统资源约束与战场实时性要求，改进现有方法在同步精度、计算复杂度和环境适应性方面的技术，"
   "以最小代价实现逼近理论最优的相参处理增益。",
   MARGIN + Inches(0.84), OBJ_Y + Inches(0.10),
   Inches(7.60), OBJ_H - Inches(0.16),
   sz=10, color=WHITE, wrap=True)
# 右侧三个关键词徽章
KW_LABELS = ["同步精度", "计算复杂度", "环境适应性"]
kw_x = MARGIN + CW - Inches(4.40)
for i, kw in enumerate(KW_LABELS):
    bx = kw_x + i * Inches(1.48)
    rect(s, bx, OBJ_Y + Inches(0.14),
         Inches(1.32), Inches(0.46),
         BADGE_COLS[i], rnd=True)
    tb(s, f"↑ {kw}", bx, OBJ_Y + Inches(0.14),
       Inches(1.32), Inches(0.46),
       sz=9, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

# ── 4. 系统需求卡片（左上）───────────────────────────────────────────
SYS_BULLETS = [
    ("• 有限通信带宽与计算资源约束下的实时相参合成", False),
    ("  分布式多节点雷达在受限条件下支撑多任务协同执行", False),
    ("• 快速变化战场态势响应能力", True),
    ("  实时处理，支撑动态战场感知与快速决策", False),
    ("• 降低节点间同步精度依赖", False),
    ("  简化系统架构，提升工程可实现性与部署灵活性", False),
]
left_card(s, LEFT_X, SYS_Y, LEFT_W, SYS_H,
          SYS_HDR, SYS_BG, SYS_ACC,
          "⚙", "系统需求",
          "System Requirements for Distributed Multi-Node Radar",
          SYS_BULLETS, hdr_h=Inches(0.42),
          title_sz=11, body_sz=8.7)

# ── 5. 性能需求卡片（左下）───────────────────────────────────────────
PER_BULLETS = [
    ("• 逼近理论最优相干处理增益与空间分集增益", True),
    ("  以更低资源消耗、更短处理时延实现最优增益", False),
    ("• 复杂电磁环境下的综合性能提升", False),
    ("  反隐身探测 · 弱小目标跟踪 · 有效抗干扰", False),
    ("• 保持高检测概率与高跟踪精度", False),
    ("  复杂干扰环境下的稳健性能保障", False),
]
left_card(s, LEFT_X, PER_Y, LEFT_W, PER_H,
          PER_HDR, PER_BG, PER_ACC,
          "📈", "性能需求",
          "Performance: Near-Optimal Coherent & Diversity Gain",
          PER_BULLETS, hdr_h=Inches(0.42),
          title_sz=11, body_sz=8.7)

# ── 6. 现有方法不足卡片（右全高）─────────────────────────────────────
GAP_ITEMS = [
    ("①", "计算复杂度过高",
     "当前方法普遍计算量大，难以在资源受限节点上实时运行"),
    ("②", "同步精度要求苛刻",
     "对节点间时空相位同步精度要求极高，工程实现代价大"),
    ("③", "非理想条件鲁棒性差",
     "在信道非理想、目标非合作等条件下性能严重下降"),
    ("④", "闭环反馈难适应高动态",
     "依赖目标参数估计闭环反馈，无法适应高机动目标场景"),
    ("⑤", "开环方法收敛受限",
     "规避目标反馈但同步实现复杂、收敛速度受限，难满足实时性"),
]
gap_card(s, RIGHT_X, BODY_Y, RIGHT_W, BODY_H, GAP_ITEMS,
         hdr_h=Inches(0.42))

# ── 7. 底部改进方向横条 ──────────────────────────────────────────────
rect(s, Inches(0), BOT_Y, SW, BOT_STRIP_H, BOT_COL)
rect(s, Inches(0), BOT_Y, Inches(0.60), BOT_STRIP_H,
     RGBColor(0x2E, 0x74, 0xB5))
tb(s, "🔧",
   Inches(0.08), BOT_Y + Inches(0.05),
   Inches(0.48), BOT_STRIP_H - Inches(0.08),
   sz=14, align=PP_ALIGN.CENTER, color=WHITE)
tb(s, "三项改进技术方向  Three Technical Improvement Directions",
   Inches(0.72), BOT_Y + Inches(0.04),
   Inches(3.60), Inches(0.20),
   sz=9.5, bold=True, color=WHITE)

IMPROV = [
    ("同步精度改进",    "低精度同步条件下鲁棒相参算法"),
    ("复杂度降低",      "轻量化实时相参合成处理架构"),
    ("环境适应性提升",  "高动态/复杂电磁环境自适应方法"),
]
mx = Inches(0.72)
m_w = (SW - mx - Inches(0.20)) / len(IMPROV)
for val, lbl in IMPROV:
    tb(s, val, mx, BOT_Y + Inches(0.24),
       m_w, Inches(0.18), sz=9, bold=True,
       color=RGBColor(0xFF, 0xE0, 0x80), align=PP_ALIGN.CENTER)
    tb(s, lbl, mx, BOT_Y + Inches(0.22 + 0.17),
       m_w, Inches(0.14), sz=7.5,
       color=RGBColor(0xC0, 0xD8, 0xFF), align=PP_ALIGN.CENTER)
    mx += m_w

# 分隔竖线
mx2 = Inches(0.72) + m_w
for _ in range(len(IMPROV) - 1):
    vline(s, mx2 - Inches(0.02), BOT_Y + Inches(0.08),
          BOT_Y + BOT_STRIP_H - Inches(0.06),
          RGBColor(0x60, 0x80, 0xC0), Pt(0.5))
    mx2 += m_w

# ── 8. 页脚 ─────────────────────────────────────────────────────────
footer(s,
       "研究目标：在资源约束与实时性前提下，突破同步精度、计算复杂度、环境适应性三项瓶颈，"
       "以最小代价获得逼近理论最优的相参处理增益，全面提升分布式雷达在复杂战场环境下的综合探测性能")

# ── 9. 保存 ─────────────────────────────────────────────────────────
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                   "coherent_synthesis_slides_v12.pptx")
prs.save(OUT)
print(f"✅  已生成：{OUT}")
