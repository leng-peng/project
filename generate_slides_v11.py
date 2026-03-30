"""
generate_slides_v11.py
──────────────────────────────────────────────────────────────────────────
多节点协同相参合成技术  ·  研究方案框架图（单页 PPT）

布局（16:9  13.33" × 7.50"）：
  ┌────────────────────────────────────────────────────────────────────┐
  │  标题栏                                                              │
  ├────────────────────────────────────────────────────────────────────┤
  │  研究背景: [三类典型协同场景] ──► [信号级相参合成方法] ──► [核心技术]  │
  ├──────────┬──────────────────────────────────────────────┬──────────┤
  │          │ 地基平台协同  │ 空中无人平台  │ 空海平台协同  │          │
  │   研究   │ 研究内容      │ 研究内容      │ 研究内容      │   核心   │
  │   范围   │               │               │               │   技术   │
  │          │     ▼         │     ▼         │     ▼         │          │
  │          │ 关键技术一    │ 关键技术二    │ 关键技术三    │          │
  ├──────────┴──────────────────────────────────────────────┴──────────┤
  │  底部横条: 五项核心技术环节                                           │
  ├────────────────────────────────────────────────────────────────────┤
  │  页脚                                                                │
  └────────────────────────────────────────────────────────────────────┘

风格：与 generate_slides_v10.py 完全一致。

运行：
    python generate_slides_v11.py
输出：
    coherent_synthesis_slides_v11.pptx
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# ━━━━━━━━━━━━━━━━━━━━━━ 调色板（与 v10 完全一致）━━━━━━━━━━━━━━━━━━━━━━
BG          = RGBColor(0xF4, 0xF7, 0xFC)
TITLE_BAR   = RGBColor(0x1F, 0x38, 0x64)
TITLE_BAR2  = RGBColor(0x2E, 0x74, 0xB5)
DARK_TEXT   = RGBColor(0x1A, 0x1A, 0x3A)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
CARD_BORDER = RGBColor(0xC8, 0xD8, 0xF0)
SEPARATOR   = RGBColor(0xC8, 0xD8, 0xF0)
FOOTER_BG   = RGBColor(0xE8, 0xF0, 0xFB)
COL_NAVY    = RGBColor(0x0D, 0x2A, 0x6E)

# 研究背景区（橙/褐色系，与 v10 一致）
BG1 = RGBColor(0xC0, 0x50, 0x00)   # 三类典型协同场景
BG2 = RGBColor(0x8B, 0x45, 0x00)   # 信号级相参合成方法
BG3 = RGBColor(0x6E, 0x25, 0x00)   # 核心技术环节

# 侧栏（深蓝紫）
SIDE_COL = RGBColor(0x0D, 0x2A, 0x6E)

# 三类协同场景 — 研究内容行颜色
SC_COLS = [
    RGBColor(0x1A, 0x5C, 0x38),   # 地基平台 — 深绿
    RGBColor(0x1F, 0x56, 0xC8),   # 空中无人 — 深蓝
    RGBColor(0xB5, 0x4A, 0x00),   # 空海平台 — 橙棕
]
SC_BGS = [
    RGBColor(0xE8, 0xF5, 0xEE),
    RGBColor(0xE8, 0xEF, 0xFE),
    RGBColor(0xFF, 0xF0, 0xE5),
]

# 关键技术行颜色
KT_COLS = [
    RGBColor(0x15, 0x7A, 0x35),   # 绿
    RGBColor(0x1F, 0x6B, 0xC4),   # 蓝
    RGBColor(0x8B, 0x2E, 0x00),   # 橙
]
KT_BGS = [
    RGBColor(0xEA, 0xF6, 0xEF),
    RGBColor(0xEB, 0xF0, 0xFF),
    RGBColor(0xFF, 0xF4, 0xEE),
]

# 底部指标栏
BOT_COL = RGBColor(0x0D, 0x2A, 0x6E)

# ━━━━━━━━━━━━━━━━━━━━━━ 尺寸常量 ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
SW      = Inches(13.33)
SH      = Inches(7.50)
MARGIN  = Inches(0.32)
TITLE_H = Inches(0.82)
CT      = TITLE_H + Inches(0.14)
FTR_H   = Inches(0.48)
FOOTER_Y = SH - FTR_H
CW      = SW - 2 * MARGIN


# ━━━━━━━━━━━━━━━━━━━━━━ 工具函数（与 v10 一致）━━━━━━━━━━━━━━━━━━━━━━━

def set_bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def rect(slide, l, t, w, h, fill, lc=None, lw=Pt(0.75), rnd=False):
    s = slide.shapes.add_shape(5 if rnd else 1, l, t, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if lc:
        s.line.color.rgb = lc
        s.line.width = lw
    else:
        s.line.fill.background()
    return s


def hline(slide, x1, y, x2, color=SEPARATOR, width=Pt(0.6)):
    c = slide.shapes.add_connector(1, x1, y, x2, y)
    c.line.color.rgb = color
    c.line.width = width


def vline(slide, x, y1, y2, color=SEPARATOR, width=Pt(0.6)):
    c = slide.shapes.add_connector(1, x, y1, x, y2)
    c.line.color.rgb = color
    c.line.width = width


def tb(slide, text, l, t, w, h,
       sz=10, bold=False, color=DARK_TEXT,
       align=PP_ALIGN.LEFT, italic=False, wrap=True):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf  = box.text_frame
    tf.word_wrap = wrap
    p   = tf.paragraphs[0]
    p.alignment = align
    r   = p.add_run()
    r.text = text
    r.font.size   = Pt(sz)
    r.font.bold   = bold
    r.font.italic = italic
    r.font.color.rgb = color
    return box


def tb_mp(slide, lines, l, t, w, h,
          sz=9, bold=False, color=DARK_TEXT,
          align=PP_ALIGN.LEFT, wrap=True):
    """Multi-paragraph text box."""
    box = slide.shapes.add_textbox(l, t, w, h)
    tf  = box.text_frame
    tf.word_wrap = wrap
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = ln
        r.font.size  = Pt(sz)
        r.font.bold  = bold
        r.font.color.rgb = color
    return box


def circle_badge(slide, cx, cy, r_inch, fill, text,
                 sz=10, bold=True, tc=WHITE):
    d = Inches(r_inch * 2)
    s = slide.shapes.add_shape(9,
                                cx - Inches(r_inch),
                                cy - Inches(r_inch), d, d)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    tf = s.text_frame
    tf.word_wrap = False
    p  = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r  = p.add_run()
    r.text = text
    r.font.size  = Pt(sz)
    r.font.bold  = bold
    r.font.color.rgb = tc


def title_bar(slide, title_main, title_sub, page, total=1):
    rect(slide, Inches(0), Inches(0), Inches(0.18), TITLE_H, COL_NAVY)
    rect(slide, Inches(0.18), Inches(0), SW - Inches(0.18), TITLE_H, TITLE_BAR)
    hline(slide, Inches(0), TITLE_H, SW, TITLE_BAR2, Pt(2.5))
    tb(slide, title_main,
       Inches(0.36), Inches(0.10), Inches(10.5), Inches(0.38),
       sz=20, bold=True, color=WHITE)
    tb(slide, title_sub,
       Inches(0.36), Inches(0.50), Inches(10.5), Inches(0.26),
       sz=10.5, color=RGBColor(0xB0, 0xCC, 0xF0))
    circle_badge(slide, SW - Inches(0.52), TITLE_H / 2, 0.28,
                 TITLE_BAR2, f"{page}/{total}", sz=11)


def footer(slide, text):
    rect(slide, Inches(0), FOOTER_Y, SW, FTR_H, FOOTER_BG)
    hline(slide, MARGIN, FOOTER_Y, SW - MARGIN, TITLE_BAR2, Pt(1.5))
    rect(slide, MARGIN, FOOTER_Y + Inches(0.10),
         Inches(0.05), Inches(0.28), TITLE_BAR2)
    tb(slide, "  " + text,
       MARGIN + Inches(0.10), FOOTER_Y + Inches(0.09),
       CW - Inches(0.14), Inches(0.34),
       sz=8.8, italic=True, color=COL_NAVY)


# ━━━━━━━━━━━━━━━━━━━━ 带标题色块的区块工具 ━━━━━━━━━━━━━━━━━━━━━━━━━━━━

def section_card(slide, l, t, w, h, hdr_color, bg_color,
                 badge, title_line1, title_line2,
                 body_lines, hdr_h=Inches(0.36),
                 title_sz=10, body_sz=8.8, border_color=None):
    """
    通用区块卡片：彩色头带 + 浅色正文区 + 左侧强调竖条
    badge      : 圆形编号标（如 "A"）
    title_line1: 粗标题
    title_line2: 副标题（可空）
    body_lines : [(text, is_accent)] list
    """
    ACC_W = Inches(0.09)
    # 阴影
    rect(slide, l + Inches(0.025), t + Inches(0.025), w, h,
         RGBColor(0xC0, 0xC8, 0xD8))
    # 主背景
    bc = border_color or hdr_color
    rect(slide, l, t, w, h, bg_color, lc=bc, lw=Pt(0.8))
    # 头带
    rect(slide, l, t, w, hdr_h, hdr_color)
    # 左侧强调竖条（正文区）
    rect(slide, l, t + hdr_h, ACC_W, h - hdr_h, hdr_color)
    # 编号圆
    if badge:
        circle_badge(slide, l + Inches(0.22), t + hdr_h / 2,
                     0.155, WHITE, badge, sz=8.5, tc=hdr_color)
    bx = l + (Inches(0.45) if badge else Inches(0.12))
    bw = w - bx + l - Inches(0.08)
    tb(slide, title_line1, bx, t + Inches(0.04),
       bw, Inches(0.22), sz=title_sz, bold=True, color=WHITE)
    if title_line2:
        tb(slide, title_line2, bx, t + hdr_h - Inches(0.14),
           bw, Inches(0.14), sz=8, color=RGBColor(0xCC, 0xEA, 0xFF))
    # 正文（每行独立 textbox，间距约 0.20"）
    pad_l = l + Inches(0.16)
    pad_w = w - Inches(0.22)
    yy    = t + hdr_h + Inches(0.06)
    lh    = Inches(0.20)          # 每行高度
    gap   = Inches(0.015)         # 行间额外间距
    for (txt, accent) in body_lines:
        clr = hdr_color if accent else DARK_TEXT
        tb(slide, txt, pad_l, yy, pad_w, lh,
           sz=body_sz, color=clr, bold=accent)
        yy += lh + gap


# ━━━━━━━━━━━━━━━━━━━━━━━━━━ MAIN ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━

prs = Presentation()
prs.slide_width  = SW
prs.slide_height = SH
blank = prs.slide_layouts[6]

s = prs.slides.add_slide(blank)
set_bg(s, BG)

# ── 1. 标题栏 ──────────────────────────────────────────────────────
title_bar(s,
          "多节点协同相参合成技术  ·  研究方案框架图",
          "Research Framework: Multi-Node Collaborative Coherent Synthesis  ·  "
          "Ground / Airborne UAV / Air-Sea Platforms",
          1)

# ── 2. 全局布局常量 ─────────────────────────────────────────────────
BOT_STRIP_H = Inches(0.38)
BOT_Y       = FOOTER_Y - BOT_STRIP_H - Inches(0.04)

TOP_Y    = CT                         # 内容区顶部
ROW_BG_H = Inches(0.60)               # 背景带高度
ROW_GAP  = Inches(0.095)              # 行间距

y_bg   = TOP_Y
y_main = y_bg + ROW_BG_H + ROW_GAP   # 主内容区起点

# 侧栏
SIDE_W   = Inches(1.12)
SIDE_GAP = Inches(0.10)
INNER_L  = MARGIN + SIDE_W + SIDE_GAP
INNER_W  = CW - 2 * (SIDE_W + SIDE_GAP)
COL_GAP  = Inches(0.10)
COL_W    = (INNER_W - 2 * COL_GAP) / 3
col_xs   = [INNER_L + i * (COL_W + COL_GAP) for i in range(3)]
RIGHT_X  = MARGIN + CW - SIDE_W

# 主内容区高度分配
TOTAL_MAIN_H = BOT_Y - y_main - Inches(0.04)
ROW_RC_H     = Inches(2.86)                         # 研究内容行
ROW_KT_H     = TOTAL_MAIN_H - ROW_RC_H - ROW_GAP   # 关键技术行

y_rc    = y_main
y_kt    = y_rc + ROW_RC_H + ROW_GAP
SIDE_H  = TOTAL_MAIN_H                              # 侧栏总高

# ── 3. 研究背景带（顶部横条）──────────────────────────────────────
BG_STEPS = [
    ("三类典型协同场景", BG1, [
        "地基平台协同场景",
        "空中无人平台协同场景",
        "空海平台协同场景",
    ]),
    ("信号级高效相参合成方法", BG2, [
        "节点内严格相参",
        "节点间正交信号发射",
        "信号级精确叠加",
    ]),
    ("五项核心技术环节", BG3, [
        "波形架构 / 多普勒处理",
        "信号级相参 / 信号级融合",
        "栅瓣控制与波位编排",
    ]),
]

n_bg      = len(BG_STEPS)
arr_w     = Inches(0.13)
bg_usable = CW - arr_w * (n_bg - 1)
bg_box_w  = bg_usable / n_bg
bx = MARGIN
for i, (lbl, col, items) in enumerate(BG_STEPS):
    rect(s, bx, y_bg, bg_box_w, ROW_BG_H, col, rnd=True)
    tb(s, lbl, bx + Inches(0.10), y_bg + Inches(0.05),
       bg_box_w - Inches(0.14), Inches(0.20),
       sz=10, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    item_y = y_bg + Inches(0.28)
    item_h = Inches(0.16)
    for it in items:
        tb(s, "• " + it, bx + Inches(0.12), item_y,
           bg_box_w - Inches(0.18), item_h,
           sz=8.2, color=WHITE)
        item_y += item_h
    if i < n_bg - 1:
        tb(s, "►", bx + bg_box_w, y_bg + (ROW_BG_H - Inches(0.20)) / 2,
           arr_w, Inches(0.20), sz=10, color=COL_NAVY,
           align=PP_ALIGN.CENTER)
    bx += bg_box_w + arr_w

# 背景带 → 研究内容列 垂直箭头
for ci in range(3):
    mid_x = col_xs[ci] + COL_W / 2
    vline(s, mid_x, y_bg + ROW_BG_H, y_rc - Inches(0.01),
          TITLE_BAR2, Pt(1.2))
    tb(s, "▼", mid_x - Inches(0.08), y_rc - Inches(0.135),
       Inches(0.16), Inches(0.14), sz=8.5,
       color=TITLE_BAR2, align=PP_ALIGN.CENTER)

# ── 4. 左侧"研究范围"侧栏 ─────────────────────────────────────────
rect(s, MARGIN + Inches(0.025), y_main + Inches(0.025),
     SIDE_W, SIDE_H, RGBColor(0xB0, 0xB8, 0xD8))
rect(s, MARGIN, y_main, SIDE_W, SIDE_H, SIDE_COL, rnd=True)
rect(s, MARGIN, y_main, Inches(0.10), SIDE_H,
     RGBColor(0x50, 0x90, 0xFF))
tb(s, "研\n究\n范\n围",
   MARGIN + Inches(0.14), y_main + (SIDE_H - Inches(0.80)) / 2,
   SIDE_W - Inches(0.16), Inches(0.80),
   sz=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
tb(s, "Research\nScope",
   MARGIN + Inches(0.06), y_main + SIDE_H - Inches(0.42),
   SIDE_W - Inches(0.08), Inches(0.40),
   sz=7, color=RGBColor(0xA0, 0xC4, 0xFF),
   align=PP_ALIGN.CENTER)

# ── 5. 右侧"核心技术"侧栏 ─────────────────────────────────────────
CORE_ITEMS = [
    "① 波形架构设计",
    "② 多普勒处理补偿",
    "③ 信号级相参叠加",
    "④ 信号级融合处理",
    "⑤ 栅瓣控制波位编排",
]
rect(s, RIGHT_X + Inches(0.025), y_main + Inches(0.025),
     SIDE_W, SIDE_H, RGBColor(0xB0, 0xB8, 0xD8))
rect(s, RIGHT_X, y_main, SIDE_W, SIDE_H, SIDE_COL, rnd=True)
rect(s, RIGHT_X + SIDE_W - Inches(0.10), y_main,
     Inches(0.10), SIDE_H, RGBColor(0x50, 0x90, 0xFF))
tb(s, "核\n心\n技\n术",
   RIGHT_X + Inches(0.02), y_main + Inches(0.04),
   SIDE_W - Inches(0.14), Inches(0.80),
   sz=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
item_y2 = y_main + Inches(0.92)
for ri in CORE_ITEMS:
    tb(s, ri, RIGHT_X + Inches(0.06), item_y2,
       SIDE_W - Inches(0.12), Inches(0.18),
       sz=7.8, color=WHITE)
    item_y2 += Inches(0.175)

# ── 6. 研究内容行（3列）──────────────────────────────────────────
RC_DATA = [
    ("地", "地基平台协同场景",
     "多栅瓣相参探测  ·  方向图栅瓣控制  ·  高效波位编排",
     [("• 多栅瓣相参探测", False),
      ("  利用多节点阵列实施相参协同探测", False),
      ("• 多节点方向图栅瓣控制", True),
      ("  精确控制各节点发射/接收方向图", False),
      ("• 栅瓣照射的高效探测波位编排方法", False),
      ("  优化多波位编排，实现高效覆盖", False),
      ("→ 地基多平台协同高效相参探测", False)]),
    ("无", "空中无人平台协同场景",
     "波形架构设计  ·  多普勒处理补偿  ·  信号级相参叠加",
     [("• 波形架构设计", False),
      ("  节点内严格相参，节点间发射正交信号", False),
      ("• 多普勒处理与补偿", False),
      ("  机动目标运动补偿，精确分离各节点回波", False),
      ("• 信号级相参叠加技术", True),
      ("  对分离后回波进行相位级精确叠加", False),
      ("→ 空中无人平台协同相参合成", False)]),
    ("海", "空海平台协同场景",
     "距离多普勒融合  ·  目标检测  ·  相干累加与增益反馈",
     [("• 距离多普勒信号级融合", False),
      ("  完成多节点信号的距离多普勒域联合处理", False),
      ("• 目标单元检测与复数提取及相位补偿", False),
      ("• 相干累加与增益计算", True),
      ("  多节点信号相干叠加，量化合成增益", False),
      ("• 增益反馈应用", False),
      ("→ 空海异构平台高效相参合成", False)]),
]

for ci, (badge, rc_title, rc_sub, items) in enumerate(RC_DATA):
    section_card(
        s, col_xs[ci], y_rc, COL_W, ROW_RC_H,
        SC_COLS[ci], SC_BGS[ci],
        badge, rc_title, rc_sub, items,
        hdr_h=Inches(0.38), title_sz=10, body_sz=8.3,
        border_color=SC_COLS[ci]
    )

# 研究内容 → 关键技术 垂直箭头
for ci in range(3):
    mid_x = col_xs[ci] + COL_W / 2
    vline(s, mid_x, y_rc + ROW_RC_H, y_kt - Inches(0.01),
          TITLE_BAR2, Pt(1.2))
    tb(s, "▼", mid_x - Inches(0.08), y_kt - Inches(0.135),
       Inches(0.16), Inches(0.14), sz=8.5,
       color=TITLE_BAR2, align=PP_ALIGN.CENTER)

# ── 7. 关键技术行（3列）─────────────────────────────────────────
KT_DATA = [
    ("①", "栅瓣控制与波位编排",
     "Grating Lobe Control & Wave Position Arrangement",
     [("• 多节点方向图栅瓣精确控制算法", False),
      ("• 高效探测波位编排优化方法", False),
      ("• 预期：多节点协同探测覆盖率提升≥30%", True)]),
    ("②", "波形架构 + 多普勒处理 + 信号级相参",
     "Waveform Design + Doppler Compensation + Signal-Level Coherent",
     [("• 节点内/间正交波形设计与严格相参参数优化", False),
      ("• 机动目标多普勒补偿与精确信号分离算法", False),
      ("• 预期：相参合成增益≥N²（误差容限内）", True)]),
    ("③", "信号级融合完整处理流程",
     "Complete Signal-Level Fusion Processing Pipeline",
     [("• 距离多普勒融合→目标检测→复数提取→相位补偿", False),
      ("• 相干累加→增益计算→增益反馈至系统优化", False),
      ("• 预期：空海平台相参合成增益≥12 dB（N=4）", True)]),
]

for ci, (badge, kt_title, kt_sub, items) in enumerate(KT_DATA):
    section_card(
        s, col_xs[ci], y_kt, COL_W, ROW_KT_H,
        KT_COLS[ci], KT_BGS[ci],
        badge, kt_title, kt_sub, items,
        hdr_h=Inches(0.34), title_sz=9.5, body_sz=8.5,
        border_color=KT_COLS[ci]
    )

# ── 8. 底部指标横栏 ──────────────────────────────────────────────
rect(s, Inches(0), BOT_Y, SW, BOT_STRIP_H, BOT_COL)
# 左侧图标区
rect(s, Inches(0), BOT_Y, Inches(0.60), BOT_STRIP_H,
     RGBColor(0x2E, 0x74, 0xB5))
tb(s, "⚙",
   Inches(0.06), BOT_Y + Inches(0.05),
   Inches(0.48), BOT_STRIP_H - Inches(0.08),
   sz=16, align=PP_ALIGN.CENTER, color=WHITE)
# 标签
tb(s, "五项核心技术环节  Five Core Technical Elements",
   Inches(0.70), BOT_Y + Inches(0.04),
   Inches(3.20), Inches(0.20),
   sz=9.5, bold=True, color=WHITE)
# 各指标
metrics = [
    ("波形架构设计",       "节点内相参·节点间正交"),
    ("多普勒处理补偿",     "机动目标精确分离回波"),
    ("信号级相参叠加",     "相位级精确叠加处理"),
    ("信号级融合处理",     "完整融合计算闭环"),
    ("栅瓣控制波位编排",   "多节点方向图优化"),
]
mx  = Inches(0.70)
m_w = (SW - mx - Inches(0.30)) / len(metrics)
for val, lbl in metrics:
    tb(s, val, mx, BOT_Y + Inches(0.24),
       m_w, Inches(0.18), sz=8.8, bold=True,
       color=RGBColor(0xFF, 0xE0, 0x80), align=PP_ALIGN.CENTER)
    tb(s, lbl, mx, BOT_Y + Inches(0.22 + 0.16),
       m_w, Inches(0.14), sz=7.5,
       color=RGBColor(0xC0, 0xD8, 0xFF), align=PP_ALIGN.CENTER)
    mx += m_w

# 分隔竖线
mx2 = Inches(0.70) + m_w
for _ in range(len(metrics) - 1):
    vline(s, mx2 - Inches(0.02), BOT_Y + Inches(0.08),
          BOT_Y + BOT_STRIP_H - Inches(0.06),
          RGBColor(0x60, 0x80, 0xC0), Pt(0.5))
    mx2 += m_w

# ── 9. 页脚 ─────────────────────────────────────────────────────
footer(s,
       "研究路径：三类典型协同场景凝练 → 多节点信号级高效相参合成方法研究 → 五项核心技术攻关 → "
       "波形架构+多普勒处理+信号级相参+信号级融合+栅瓣控制 全链路闭环，实现多平台高效协同相参合成")

# ── 10. 保存 ─────────────────────────────────────────────────────
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                   "coherent_synthesis_slides_v11.pptx")
prs.save(OUT)
print(f"✅  已生成：{OUT}")
